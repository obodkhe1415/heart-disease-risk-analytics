import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import pandas as pd
from sklearn.model_selection import train_test_split, cross_val_score
from sklearn.ensemble import RandomForestRegressor
from sklearn.metrics import r2_score, mean_absolute_error

# Compute real model metrics from the dataset
_df = pd.read_csv(os.path.join(os.path.dirname(__file__), 'data', 'heart_disease_data.csv'))
_X = _df[['age_years','resting_bp_mmhg','cholesterol_mgdl','fasting_bs_mgdl','max_heart_rate','exercise_angina','st_depression']]
_y = _df['heart_disease_risk_pct']
_Xtr, _Xte, _ytr, _yte = train_test_split(_X, _y, test_size=0.2, random_state=42)
_m = RandomForestRegressor(n_estimators=100, random_state=42).fit(_Xtr, _ytr)
R2 = r2_score(_yte, _m.predict(_Xte))
MAE = mean_absolute_error(_yte, _m.predict(_Xte))
CV_R2 = cross_val_score(RandomForestRegressor(n_estimators=100, random_state=42), _X, _y, cv=5, scoring='r2').mean()

doc = docx.Document()

# Title Section
title_p = doc.add_paragraph()
title_run = title_p.add_run('Executive Project Report: Heart Disease & Cardiovascular Risk Analytics Platform')
title_run.font.size = Pt(22)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(11, 60, 93)

sub_p = doc.add_paragraph()
sub_run = sub_p.add_run('AICTE | IBM SkillsBuild Data Analytics with AI Academic Internship 2026\n')
sub_run.font.size = Pt(13)
sub_run.font.italic = True
sub_run.font.color.rgb = RGBColor(100, 100, 100)

doc.add_paragraph('Author: Student Name | Internship Track: Data Analytics with AI')
doc.add_paragraph('Organization: BharatCares (SMEC Trust) & IBM SkillsBuild')

# Section 1: Executive Summary
doc.add_heading('1. Executive Summary', level=1)
doc.add_paragraph(
    'This project delivers an end-to-end Machine Learning and Business Intelligence solution for Healthcare Analytics & Cardiovascular Risk Assessment. The application predicts patient cardiovascular risk scores based on clinical biomarkers (blood pressure, cholesterol, fasting sugar, max heart rate, and ST depression) and provides actionable clinical recommendations for healthcare leadership.'
)

# Section 2: Business Intelligence Dashboard Hierarchy
doc.add_heading('2. The 5-Level Dashboard Hierarchy', level=1)
p_hier = doc.add_paragraph()
p_hier.add_run('• Level 1 (KPIs): ').bold = True
p_hier.add_run('Evaluated patient count, high-risk patient rate %, average resting BP (mmHg), and model R² accuracy score.\n')
p_hier.add_run('• Level 2 (Trends): ').bold = True
p_hier.add_run('Scatter plot of resting BP vs cholesterol with OLS trendlines broken down by chest pain type.\n')
p_hier.add_run('• Level 3 (Drivers): ').bold = True
p_hier.add_run('Feature correlation analysis quantifying the impact of age, BP, cholesterol, and ST depression on risk score.\n')
p_hier.add_run('• Level 4 (Risks & Opportunities): ').bold = True
p_hier.add_run('Interactive AI cardiovascular risk predictor and clinical risk stratification engine.\n')
p_hier.add_run('• Level 5 (Recommended Actions): ').bold = True
p_hier.add_run('Preventive cardiology referral guidelines, automated EHR clinical alerts, and wellness intervention programs.')

# Section 3: Technology Stack Table
doc.add_heading('3. Technology Stack', level=1)
table = doc.add_table(rows=5, cols=2)
table.style = 'Table Grid'

hdr = table.rows[0].cells
hdr[0].text = 'Component / Layer'
hdr[1].text = 'Technology Used'

tech_data = [
    ('Data Engineering & Clinical Analytics', 'Pandas, NumPy'),
    ('Machine Learning Algorithm', 'Scikit-Learn (Random Forest Regressor)'),
    ('Interactive Dashboard UI', 'Streamlit, Plotly Express'),
    ('Automated Reporting', 'Python-Docx')
]

for idx, (layer, tech) in enumerate(tech_data, start=1):
    row = table.rows[idx].cells
    row[0].text = layer
    row[1].text = tech

# Section 4: Model Evaluation
doc.add_heading('4. Machine Learning Model Performance', level=1)
doc.add_paragraph(
    'The Random Forest model was trained on clinical cardiovascular biomarkers. Key metrics:'
)
p_metrics = doc.add_paragraph()
p_metrics.add_run('• R² Score (Risk Forecast Accuracy): ').bold = True
p_metrics.add_run(f'{R2:.2f} (hold-out test set, n={len(_yte)}); {CV_R2:.2f} (5-fold cross-validation)\n')
p_metrics.add_run('• Mean Absolute Error (MAE): ').bold = True
p_metrics.add_run(f'{MAE:.1f} risk-score points (hold-out test set)\n')

# Section 5: Strategic Executive Recommendations
doc.add_heading('5. Strategic Executive Recommendations', level=1)
doc.add_paragraph('1. Mandate preventive cardiology consultations for patients with a predicted risk score > 50%.')
doc.add_paragraph('2. Deploy automated EHR clinical alerts for serum cholesterol levels exceeding 240 mg/dL.')
doc.add_paragraph('3. Establish community cardiovascular wellness clinics targeting individuals aged 50+.')

os.makedirs('reports', exist_ok=True)
output_filename = os.path.join('reports', 'HeartDisease_ProjectReport.docx')
doc.save(output_filename)
print(f"Project report generated successfully: {output_filename}")
